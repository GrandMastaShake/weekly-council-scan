"""Build the uploadable parts of Dad's Claude kit into dad-kit/dist/.

    python dad-kit/build_kit.py

Produces:
  dist/stock-research.zip, dist/golf-course-ops.zip
      One per skill. claude.ai wants the skill folder as the zip's root and
      the folder name equal to the skill's `name`, so both are checked
      before anything is written.
  dist/Investing Journal.xlsx, dist/League Manager.xlsx,
  dist/Golf Course Organizer.xlsx
      Blank starter workbooks. Needs openpyxl (pip install openpyxl); the
      zips do not, and are still built without it.

dist/ is not committed: everything in it is derived from the files beside
this script, so a committed copy would only be a second copy that drifts.
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

KIT = Path(__file__).resolve().parent
SKILLS = KIT / "skills"
DIST = KIT / "dist"

NAME_RE = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")
NAME_MAX = 64
# claude.ai's upload form caps the description at 200 characters, tighter
# than the 1024 the API allows. The tighter limit is the one that bites.
DESCRIPTION_MAX = 200


def read_frontmatter(skill_md: Path) -> dict[str, str]:
    """The flat `key: value` frontmatter of a SKILL.md. No YAML dependency:
    the kit's frontmatter is two plain scalar lines and must stay that way,
    so anything fancier is itself an error."""
    text = skill_md.read_text(encoding="utf-8")
    lines = text.splitlines()
    if not lines or lines[0].strip() != "---":
        raise ValueError(f"{skill_md}: does not open with a '---' frontmatter line")
    fields: dict[str, str] = {}
    for line in lines[1:]:
        if line.strip() == "---":
            return fields
        key, sep, value = line.partition(":")
        if not sep or not key.strip():
            raise ValueError(f"{skill_md}: frontmatter line is not 'key: value': {line!r}")
        fields[key.strip()] = value.strip()
    raise ValueError(f"{skill_md}: frontmatter is never closed with '---'")


def check_skill(skill_dir: Path) -> list[str]:
    """Every reason claude.ai would reject this skill, or the kit's own rules
    would. Empty means it is good to zip."""
    problems: list[str] = []
    skill_md = skill_dir / "SKILL.md"
    if not skill_md.is_file():
        return [f"{skill_dir.name}: no SKILL.md"]
    try:
        fm = read_frontmatter(skill_md)
    except ValueError as exc:
        return [str(exc)]

    name = fm.get("name", "")
    desc = fm.get("description", "")
    if not NAME_RE.match(name) or len(name) > NAME_MAX:
        problems.append(f"{skill_dir.name}: name {name!r} must be lowercase words joined by hyphens, <= {NAME_MAX} chars")
    if name != skill_dir.name:
        problems.append(f"{skill_dir.name}: name {name!r} does not match its folder")
    if not desc:
        problems.append(f"{skill_dir.name}: no description")
    elif len(desc) > DESCRIPTION_MAX:
        problems.append(f"{skill_dir.name}: description is {len(desc)} chars, limit {DESCRIPTION_MAX}")
    if "<" in desc or ">" in desc:
        problems.append(f"{skill_dir.name}: description contains angle brackets")

    body = skill_md.read_text(encoding="utf-8")
    for ref in sorted(set(re.findall(r"`(references/[^`]+\.md)`", body))):
        if not (skill_dir / ref).is_file():
            problems.append(f"{skill_dir.name}: SKILL.md points at {ref}, which does not exist")
    return problems


def skill_dirs() -> list[Path]:
    return sorted(p for p in SKILLS.iterdir() if p.is_dir() and not p.name.startswith("."))


def zip_skill(skill_dir: Path, out_dir: Path) -> Path:
    """Zip one skill with its folder as the archive root. Forward slashes and
    a sorted file order, so the same tree always gives the same archive."""
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{skill_dir.name}.zip"
    files = sorted(p for p in skill_dir.rglob("*") if p.is_file() and "__pycache__" not in p.parts)
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in files:
            arcname = f"{skill_dir.name}/{path.relative_to(skill_dir).as_posix()}"
            info = zipfile.ZipInfo(arcname, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, path.read_bytes())
    return target


# --- workbooks -------------------------------------------------------------

STATUS = ["Open", "In progress", "Waiting", "Done"]
PRIORITY = ["1", "2", "3"]
AREAS = [
    "tee", "fairway", "rough", "green", "bunker", "cart path", "practice area",
    "clubhouse", "pro shop", "cart barn", "maintenance shop", "parking lot", "entrance",
]

GOLF_SHEETS = {
    "Tasks": (
        ["ID", "Task", "Hole", "Area", "Who", "Priority", "Due", "Status", "Done on", "Notes"],
        [6, 44, 7, 16, 16, 9, 12, 12, 12, 40],
        {"Area": AREAS, "Priority": PRIORITY, "Status": STATUS},
    ),
    "Schedule": (
        ["Date", "Day", "Shift", "Start", "End", "Role", "Name", "Confirmed", "Notes"],
        [12, 11, 12, 9, 9, 18, 20, 11, 36],
        {"Confirmed": ["Yes", "No"]},
    ),
    "Events": (
        ["Event", "Date", "Day", "Item", "Who", "Due", "Status", "Cost", "Notes"],
        [26, 12, 11, 36, 16, 12, 12, 11, 36],
        {"Status": STATUS},
    ),
    "Sponsors": (
        ["Event", "Sponsor", "Level", "Amount", "Paid", "Logo received", "Sign ordered", "Thank-you sent", "Notes"],
        [24, 26, 12, 11, 8, 14, 13, 15, 30],
        {"Paid": ["Yes", "No"], "Logo received": ["Yes", "No"], "Sign ordered": ["Yes", "No"], "Thank-you sent": ["Yes", "No"]},
    ),
    "Maintenance Log": (
        ["Date", "Hole", "Area", "What was done", "Who", "Hours", "Materials used", "Next due", "Notes"],
        [12, 7, 16, 40, 16, 8, 26, 12, 36],
        {"Area": AREAS},
    ),
    "Equipment": (
        ["Machine", "Hour meter", "Service done", "Date", "Next due", "Who", "Notes"],
        [26, 11, 36, 12, 16, 16, 36],
        {},
    ),
    "Inventory": (
        ["Item", "Where kept", "On hand", "Reorder at", "Supplier", "Last ordered", "Notes"],
        [30, 18, 9, 11, 22, 13, 36],
        {},
    ),
    "Contacts": (
        ["Name", "Role", "Group", "Phone", "Email", "Best way to reach", "Notes"],
        [24, 22, 18, 15, 28, 16, 30],
        {},
    ),
}

GOLF_HOWTO = [
    "Golf Course Organizer",
    "",
    "One workbook for the course: tasks, schedules, events, sponsors, maintenance, equipment, supplies and contacts.",
    "",
    "How to use it with Claude:",
    "1. Keep this file on your computer (or in Google Drive / OneDrive).",
    "2. In the Golf project, start a chat, attach this file, and say what changed.",
    "   Example: 'Add these jobs from today' or 'Build the volunteer schedule for the scramble'.",
    "3. Claude hands back the whole updated file. Save it over the old one.",
    "",
    "Rules that keep it useful:",
    "- Never delete a finished task. Set Status to Done and fill in Done on. Next year, this year's list is the plan.",
    "- Priority: 1 = safety or today, 2 = this week, 3 = when there is time.",
    "- Hole + Area together say where a job is, e.g. Hole 7, Area green.",
    "- Blank is better than a guess. Claude will ask for what is missing.",
]

JOURNAL_SHEETS = {
    "Holdings": (
        ["Ticker", "Name", "Account type", "Shares", "Avg cost", "Cost basis", "Last price", "Price date",
         "Price source", "Value", "Gain $", "Gain %", "% of portfolio", "Sector", "Plan on file?"],
        [9, 24, 15, 9, 10, 12, 10, 12, 16, 12, 11, 9, 14, 18, 13],
        {"Plan on file?": ["Yes", "No"]},
    ),
    "Trade Plans": (
        ["Date", "Ticker", "Action", "Why (one sentence)", "Evidence (dated)", "What would prove me wrong",
         "Stop", "Stop is a real order?", "Target or review date", "Size % of portfolio", "Dollars at risk",
         "Events before review", "Same bet as", "Outcome", "Lesson"],
        [12, 9, 9, 40, 40, 34, 9, 12, 14, 12, 12, 26, 16, 26, 34],
        {"Action": ["Buy", "Add", "Trim", "Sell"], "Stop is a real order?": ["Yes", "No"]},
    ),
    "Watchlist": (
        ["Ticker", "Why I'm watching", "Buy zone", "Trigger", "Added", "Last checked", "Source", "Still valid?"],
        [9, 40, 12, 30, 12, 13, 26, 11],
        {"Still valid?": ["Yes", "No"]},
    ),
    "Limit Orders": (
        ["Ticker", "Side", "Limit price", "Placed", "Times re-confirmed", "Last re-confirmed",
         "Why it is still right", "Status"],
        [9, 7, 11, 12, 12, 14, 44, 12],
        {"Side": ["Buy", "Sell"], "Status": ["Active", "Filled", "Cancelled"]},
    ),
    "Weekly Review": (
        ["Week of", "Market in one line", "Every holding's reason still true?", "Actions taken",
         "Limit orders (kept / re-priced / cancelled)", "Next week's events", "Notes"],
        [12, 40, 22, 30, 26, 34, 34],
        {},
    ),
}

JOURNAL_HOWTO = [
    "Investing Journal",
    "",
    "Your holdings, your plans, your watchlist and your weekly review in one place.",
    "",
    "How to use it with Claude:",
    "1. Keep this file on your computer. Attach it to a chat in the Stocks project when you want it updated.",
    "   Examples: 'Update my holdings with Friday's closes', 'Help me write a plan for buying X', 'Let's do my weekly review'.",
    "2. Claude hands back the whole updated file. Save it over the old one.",
    "",
    "The habits this is built around (from the family Council's own logs):",
    "- Every buy gets a plan first: why, what would prove you wrong, the stop, the size.",
    "  If you can't say what would prove you wrong, the plan isn't ready.",
    "- Every price gets a date and a source. A price with no date is a rumor.",
    "- A limit order re-confirmed 3 weeks running with no fill: decide if you want the stock at market, or cancel it.",
    "- Write a Weekly Review row even when nothing changed. 'Nothing to do' and 'nobody looked' look the same later.",
    "- Two holdings that would be wrong for the same reason are one bet.",
    "",
    "Cash: add it as a Holdings row (Ticker CASH, Shares = dollars, Avg cost 1, Last price 1) so % of portfolio counts it.",
    "",
    "No account numbers or passwords in this file. Claude doesn't need them.",
]


LEAGUE_SHEETS = {
    "Roster": (
        ["Name", "Team", "Phone", "Email", "Handicap", "Handicap as of", "Dues paid", "Paid on", "Active", "Notes"],
        [22, 14, 15, 26, 10, 14, 10, 12, 8, 30],
        {"Dues paid": ["Yes", "No"], "Active": ["Yes", "No"]},
    ),
    "Schedule": (
        ["Week", "Date", "Day", "Start", "Nine", "Match or group", "Side A", "Side B", "Status", "Makeup date", "Notes"],
        [7, 12, 11, 9, 8, 14, 20, 20, 12, 13, 30],
        {"Nine": ["Front", "Back", "18"], "Status": ["Scheduled", "Played", "Rained out", "Makeup"]},
    ),
    "Weekly Scores": (
        ["Week", "Date", "Player", "Team", "Gross", "Handicap used", "Net", "Points", "Skins won", "Sub for", "Notes"],
        [7, 12, 22, 14, 8, 10, 8, 8, 10, 18, 34],
        {},
    ),
    "Standings": (
        ["Player", "Team", "Points", "Rounds", "Avg net", "Rank"],
        [22, 14, 9, 9, 9, 7],
        {},
    ),
    "Team Standings": (
        ["Team", "Points", "Rank"],
        [20, 9, 7],
        {},
    ),
    "Subs": (
        ["Name", "Phone", "Email", "Handicap", "Nights available", "Times subbed", "Last subbed", "Notes"],
        [22, 15, 26, 10, 18, 12, 12, 30],
        {},
    ),
    "Money": (
        ["Date", "Who", "What", "In", "Out", "Balance", "Notes"],
        [12, 22, 12, 10, 10, 11, 34],
        {"What": ["Dues", "Skins", "Prize", "Payout", "Expense", "Other"]},
    ),
}

LEAGUE_HOWTO = [
    "League Manager",
    "",
    "One workbook for the season: roster, schedule, weekly scores, standings, subs and money.",
    "",
    "How to use it with Claude:",
    "1. Keep this file on your computer. After league night, start a chat in the Golf project,",
    "   attach this file and photos of the scorecards, and say 'Enter week 5'.",
    "2. Claude reads each card, checks the holes add up to the total written on it, and shows you",
    "   the scores to confirm before doing anything else.",
    "3. It fills in Weekly Scores; Standings and Team Standings add themselves up.",
    "   It hands back the whole file. Save it over the old one.",
    "",
    "Rules that keep it right:",
    "- One row per player per week in Weekly Scores. Net = Gross minus Handicap used, filled in for you.",
    "- Nobody gets a made-up score. A player who did not play gets what the league's absence rule says.",
    "- Never delete a rained-out week from Schedule. Set its Status to Rained out.",
    "- Standings lists everyone on the Roster, in Roster order, ranked by points (most first).",
    "  Tie-breakers come from your league rules; Claude applies them when you ask for final standings.",
    "- Subs go on the Subs sheet. Times subbed counts itself from Weekly Scores.",
    "- Money: one row for every dollar in or out. Balance is a running total and should match the cash box.",
]


def _league_formulas(title: str, ws, players: int = 60, teams: int = 30, rows: int = 600) -> None:
    """Standings that add themselves up from Weekly Scores, so entering a
    week is one sheet of typing and nobody re-totals points by hand. Only
    functions every Excel since 2010 and LibreOffice know: no SORT, no
    MAXIFS, which would show #NAME? on an older copy of Office."""
    scores = "'Weekly Scores'"
    if title == "Weekly Scores":
        for r in range(2, rows + 2):
            ws[f"G{r}"] = f'=IF(AND(E{r}<>"",F{r}<>""),E{r}-F{r},"")'
    elif title == "Standings":
        last = players + 1
        for r in range(2, last + 1):
            ws[f"A{r}"] = f'=IF(Roster!A{r}="","",Roster!A{r})'
            ws[f"B{r}"] = f'=IF(A{r}="","",IF(Roster!B{r}="","",Roster!B{r}))'
            ws[f"C{r}"] = f'=IF(A{r}="","",SUMIF({scores}!$C:$C,A{r},{scores}!$H:$H))'
            ws[f"D{r}"] = f'=IF(A{r}="","",COUNTIF({scores}!$C:$C,A{r}))'
            ws[f"E{r}"] = f'=IF(A{r}="","",IFERROR(AVERAGEIF({scores}!$C:$C,A{r},{scores}!$G:$G),""))'
            ws[f"F{r}"] = f'=IF(OR(A{r}="",D{r}=0),"",RANK(C{r},$C$2:$C${last}))'
            ws[f"E{r}"].number_format = "0.0"
    elif title == "Team Standings":
        last = teams + 1
        for r in range(2, last + 1):
            ws[f"B{r}"] = f'=IF(A{r}="","",SUMIF({scores}!$D:$D,A{r},{scores}!$H:$H))'
            ws[f"C{r}"] = f'=IF(A{r}="","",RANK(B{r},$B$2:$B${last}))'
    elif title == "Subs":
        for r in range(2, players + 2):
            ws[f"F{r}"] = f'=IF(A{r}="","",COUNTIF({scores}!$C:$C,A{r}))'
    elif title == "Money":
        for r in range(2, rows + 2):
            ws[f"F{r}"] = f'=IF(AND(D{r}="",E{r}=""),"",SUM($D$2:D{r})-SUM($E$2:E{r}))'
            for col in "DEF":
                ws[f"{col}{r}"].number_format = "$#,##0.00"


def _holdings_formulas(ws, rows: int = 60) -> None:
    """Cost basis, value, gain and weight computed in the sheet, so a pasted
    price updates everything and nobody re-does the arithmetic by hand."""
    for r in range(2, rows + 2):
        ws[f"F{r}"] = f'=IF(AND(D{r}<>"",E{r}<>""),D{r}*E{r},"")'
        ws[f"J{r}"] = f'=IF(AND(D{r}<>"",G{r}<>""),D{r}*G{r},"")'
        ws[f"K{r}"] = f'=IF(AND(J{r}<>"",F{r}<>""),J{r}-F{r},"")'
        ws[f"L{r}"] = f'=IF(AND(K{r}<>"",F{r}<>"",F{r}<>0),K{r}/F{r},"")'
        ws[f"M{r}"] = f'=IF(AND(J{r}<>"",SUM($J$2:$J${rows + 1})<>0),J{r}/SUM($J$2:$J${rows + 1}),"")'
        for col, fmt in (("E", "$#,##0.00"), ("F", "$#,##0.00"), ("G", "$#,##0.00"),
                         ("J", "$#,##0.00"), ("K", "$#,##0.00"), ("L", "0.0%"), ("M", "0.0%")):
            ws[f"{col}{r}"].number_format = fmt


def build_workbook(path: Path, howto: list[str], sheets: dict, after=None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "How To Use"
    ws.column_dimensions["A"].width = 120
    for i, line in enumerate(howto, start=1):
        cell = ws.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True, size=16)

    header_fill = PatternFill("solid", fgColor="1F4E3D")
    header_font = Font(bold=True, color="FFFFFF")
    for title, (headers, widths, choices) in sheets.items():
        ws = wb.create_sheet(title)
        for c, (head, width) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=1, column=c, value=head)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        for head, options in choices.items():
            col = get_column_letter(headers.index(head) + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}1000")
        if after:
            after(title, ws)
    wb.save(path)
    return path


def build_workbooks(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)

    def journal_extras(title, ws):
        if title == "Holdings":
            _holdings_formulas(ws)

    return [
        build_workbook(out_dir / "Investing Journal.xlsx", JOURNAL_HOWTO, JOURNAL_SHEETS, journal_extras),
        build_workbook(out_dir / "League Manager.xlsx", LEAGUE_HOWTO, LEAGUE_SHEETS, _league_formulas),
        build_workbook(out_dir / "Golf Course Organizer.xlsx", GOLF_HOWTO, GOLF_SHEETS),
    ]


def main() -> int:
    problems = [p for d in skill_dirs() for p in check_skill(d)]
    if problems:
        print("Not building; fix these first:")
        for p in problems:
            print(f"  - {p}")
        return 1

    for d in skill_dirs():
        print(f"built {zip_skill(d, DIST).relative_to(KIT.parent)}")

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("skipped the workbooks: openpyxl is not installed (pip install openpyxl)")
        return 0
    for path in build_workbooks(DIST):
        print(f"built {path.relative_to(KIT.parent)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
